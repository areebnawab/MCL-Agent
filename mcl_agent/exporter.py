"""Create a formatted Excel workbook from agent results."""

from __future__ import annotations

from io import BytesIO
from typing import Any

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


def create_results_workbook(rows: list[dict[str, Any]], summary: dict[str, int]) -> bytes:
    """Export a summary and evidence table, returning downloadable XLSX bytes."""
    output = BytesIO()
    result_frame = pd.DataFrame(rows)
    if result_frame.empty:
        result_frame = pd.DataFrame(columns=["REVIEW STATUS", "SOURCE DOCUMENT", "SOURCE EVIDENCE"])
    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        pd.DataFrame([{"Measure": key.replace("_", " ").title(), "Count": value} for key, value in summary.items()]).to_excel(writer, sheet_name="Summary", index=False)
        result_frame.to_excel(writer, sheet_name="Renewal Results", index=False)
        workbook = writer.book
        for sheet in workbook.worksheets:
            sheet.freeze_panes = "A2"
            sheet.auto_filter.ref = sheet.dimensions
            for cell in sheet[1]:
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = PatternFill("solid", fgColor="17365D")
                cell.alignment = Alignment(wrap_text=True, vertical="center")
            sheet.row_dimensions[1].height = 32
            for column_cells in sheet.columns:
                width = min(max(max(len(str(cell.value or "")) for cell in column_cells) + 2, 12), 48)
                sheet.column_dimensions[get_column_letter(column_cells[0].column)].width = width
            if sheet.title == "Renewal Results":
                for row_index in range(2, sheet.max_row + 1):
                    status_cell = sheet.cell(row_index, next((cell.column for cell in sheet[1] if cell.value == "REVIEW STATUS"), 1))
                    if str(status_cell.value).startswith("Review"):
                        status_cell.fill = PatternFill("solid", fgColor="FCE4D6")
    return output.getvalue()
