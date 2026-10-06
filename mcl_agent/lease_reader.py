"""Extract searchable text from lease documents stored on disk or uploaded."""

from pathlib import Path

from docx import Document
from pypdf import PdfReader

from .config import SUPPORTED_LEASE_EXTENSIONS
from .models import LeaseDocument


def _read_one_document(file_path: Path) -> LeaseDocument:
    """Read one supported document and keep page labels beside extracted text."""
    try:
        if file_path.suffix.lower() == ".pdf":
            reader = PdfReader(str(file_path))
            # Layout mode retains table columns such as product, quantity, and term values.
            pages = [(f"page {index + 1}", page.extract_text(extraction_mode="layout") or "") for index, page in enumerate(reader.pages)]
            # Preserve page boundaries so the agent can cite the source page for extracted evidence.
            text = "\n\n".join(f"[{label}]\n{page_text}" for label, page_text in pages)
            return LeaseDocument(filename=file_path.name, text=text, page_labels=[label for label, _ in pages])
        if file_path.suffix.lower() == ".docx":
            document = Document(str(file_path))
            text = "\n".join(paragraph.text for paragraph in document.paragraphs if paragraph.text.strip())
            for table in document.tables:
                text += "\n" + "\n".join(" | ".join(cell.text for cell in row.cells) for row in table.rows)
            return LeaseDocument(filename=file_path.name, text=text, page_labels=["document"])
        return LeaseDocument(filename=file_path.name, text=file_path.read_text(encoding="utf-8", errors="replace"), page_labels=["document"])
    except Exception as exc:
        return LeaseDocument(filename=file_path.name, text="", error=str(exc))


def load_lease_documents(folder: Path, uploaded_paths: list[Path] | None = None) -> list[LeaseDocument]:
    """Read supported lease files, returning an empty list when none are provided."""
    paths: dict[str, Path] = {}
    if folder.exists() and folder.is_dir():
        for path in folder.rglob("*"):
            if path.is_file() and path.suffix.lower() in SUPPORTED_LEASE_EXTENSIONS:
                paths[str(path.resolve())] = path
    for path in uploaded_paths or []:
        if path.suffix.lower() in SUPPORTED_LEASE_EXTENSIONS:
            paths[str(path.resolve())] = path
    return [_read_one_document(path) for path in sorted(paths.values(), key=lambda item: item.name.lower())]
